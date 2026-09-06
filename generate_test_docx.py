import os
import io
from PIL import Image, ImageDraw, ImageFont
import docx
from docx.shared import Inches

def generate_test_docx(output_path: str):
    # 1. Generate an image containing revenue & customer metrics
    img = Image.new('RGB', (600, 220), color=(245, 247, 250))
    d = ImageDraw.Draw(img)
    d.rectangle([(10, 10), (590, 210)], outline=(30, 80, 160), width=3)
    d.text((30, 30), "QUARTERLY PERFORMANCE REPORT - METRICS", fill=(10, 30, 80))
    d.text((30, 80), "Total Revenue: $4,250,000", fill=(0, 100, 40))
    d.text((30, 130), "Active Enterprise Customers: 1,840", fill=(150, 40, 20))
    d.text((30, 175), "Target Q4 Growth Rate: +18.5%", fill=(80, 80, 80))
    
    img_buffer = io.BytesIO()
    img.save(img_buffer, format="PNG")
    img_buffer.seek(0)

    # 2. Create the docx
    doc = docx.Document()
    doc.add_heading("Quarterly Executive Summary & Engineering Task", level=1)
    
    p1 = doc.add_paragraph("This document contains financial performance metrics from the recent audit alongside an urgent software requirement.")
    
    doc.add_heading("Financial Dashboard Screenshot", level=2)
    doc.add_picture(img_buffer, width=Inches(4.5))
    
    doc.add_heading("Engineering Requirements", level=2)
    p2 = doc.add_paragraph(
        "Software Requirement:\n"
        "Write a Python function `safe_api_transform(response_data)` that takes an API payload. "
        "If `response_data` is None, or if the 'items' key inside `response_data` is None, it must safely return an empty list `[]`. "
        "Otherwise, return the items list. Include unit test docstrings and type hints."
    )

    doc.save(output_path)
    print(f"Created test DOCX at: {output_path}")

if __name__ == "__main__":
    import config
    out_dir = config.TEMP_DIR
    os.makedirs(out_dir, exist_ok=True)
    generate_test_docx(os.path.join(out_dir, "test_sample_report.docx"))
